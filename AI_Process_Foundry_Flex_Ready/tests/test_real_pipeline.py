"""Unit tests + real local media/export checks. No substituted Foundry responses."""
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engine.config import Settings, https_endpoint
from engine.media import plan_segments
from engine.storage import LocalStore, read_json
from engine.transcription import PendingAuthorization, check_mai_authorization, normalize_mai
from engine.vision import Analysis, parse_model_json, related_text
from real_pipeline import run_pipeline
from transcript_ingestion import read_transcript
from poc_core import presentation

ROOT = Path(__file__).resolve().parents[1]


class PipelineTests(unittest.TestCase):
    def test_complete_timeline(self):
        for duration in (300, 3600, 18000, 18000.25):
            segments = plan_segments(duration, 120)
            self.assertEqual(segments[0]['start'], 0)
            self.assertEqual(segments[-1]['end'], duration)
            self.assertEqual(sum(s['end'] - s['start'] for s in segments), duration)
            self.assertTrue(all(a['end'] == b['start'] for a, b in zip(segments, segments[1:])))

    def test_no_default_deployment(self):
        with patch.dict(os.environ, {'AZURE_VISION_DEPLOYMENT': ''}):
            with self.assertRaises(ValueError):
                Settings.from_env()
        self.assertFalse(Settings(use_ai=False).use_ai)
        with self.assertRaises(ValueError):
            https_endpoint('https://resource/openai/deployments/old', v1=True)

    def test_pending_mai_never_invokes(self):
        with patch.dict(os.environ, {'MAI_AUTHORIZED': 'false'}):
            with self.assertRaises(PendingAuthorization):
                check_mai_authorization()

    def test_strict_schema_and_incomplete_contract(self):
        self.assertEqual(parse_model_json('{"steps":[]}'), {'steps': []})
        with self.assertRaises(ValueError):
            Analysis.model_validate({'steps': [{'frame_index': 1}]})
        with self.assertRaises(ValueError):
            normalize_mai({'text': 'no timestamps'}, 0, 120, 'segment_0')
        # Contract-only fixture: deliberately not a live service validation.
        rows = normalize_mai({'phrases': [{'offsetMilliseconds': 500, 'durationMilliseconds': 1000,
                              'text': 'Texto de prueba', 'speaker': 0}]}, 120, 120, 'segment_1')
        self.assertEqual(rows[0]['start'], 120.5)
        self.assertEqual(rows[0]['speaker'], 'segment_1:0')

    def test_filter_and_no_text_slicing(self):
        segments = [{'start': 2, 'end': 4, 'speaker': 'Ana', 'text': 'Pregunta'},
                    {'start': 3, 'end': 5, 'speaker': 'Profe', 'text': 'x' * 7000}]
        text = related_text(segments, 4, main_speaker='Profe')
        self.assertNotIn('Pregunta', text)
        self.assertIn('x' * 7000, text)

    def test_vtt_srt_json(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            samples = {'vtt': 'WEBVTT\n\n1\n00:00:00.000 --> 00:00:01.000\n<v Ana>Crear</v>\n',
                       'srt': '1\n00:00:00,000 --> 00:00:01,000\nCrear\n',
                       'json': json.dumps({'segments': [{'start': 0, 'end': 1, 'text': 'Crear'}]})}
            for extension, content in samples.items():
                file = root / ('transcript.' + extension)
                file.write_text(content)
                self.assertEqual(read_transcript(file)[0]['text'], 'Crear')
            invalid = root / 'invalid.json'
            invalid.write_text('{"segments":[{"text":"no time"}]}')
            with self.assertRaises(ValueError):
                read_transcript(invalid)

    def test_legacy_truncation_is_rejected(self):
        with self.assertRaises(ValueError):
            run_pipeline('not-opened.mp4', use_ai=False, limit_minutes=6)

    def test_real_extraction_resume_and_no_fabricated_documents(self):
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ, {
                'PIPELINE_DATA_DIR': folder, 'USE_AZURE_BLOB': 'false'}):
            result = run_pipeline(str(ROOT / 'assets/video_demo_ficticio.mp4'), use_ai=False,
                                  segment_seconds=3, frame_every_seconds=1)
            self.assertEqual(result['status'], 'extracted_without_ai', result.get('error'))
            manifest = json.loads(Path(result['manifest']).read_text())
            self.assertEqual(manifest['processed_seconds'], manifest['video_duration_seconds'])
            self.assertGreater(len(manifest['coverage']), 1)
            self.assertTrue(all(c['frames'] > 0 for c in manifest['coverage']))
            self.assertIsNone(result['pptx'])
            before = Path(result['job']).read_bytes()
            resumed = run_pipeline(None, resume=result['project_id'])
            self.assertEqual(resumed['status'], result['status'])
            self.assertEqual(Path(result['job']).read_bytes(), before)

    def test_failed_media_unit_can_resume_without_losing_prior_segments(self):
        from engine.jobs import create_job, advance_job, TERMINAL
        with tempfile.TemporaryDirectory() as directory:
            store = LocalStore(directory)
            store.upload('raw', 'input.mp4', ROOT / 'assets/video_demo_ficticio.mp4')
            state = create_job(store, 'input.mp4', settings=Settings(use_ai=False, segment_seconds=6,
                                                                   frame_every_seconds=1))
            job_id = state['project_id']
            state = advance_job(store, job_id)
            self.assertEqual(state['next_segment'], 1)
            checkpoint = store.get('analysis', f'{job_id}/segment_000000/result.json')
            # Fault injection only, never a fabricated successful AI response.
            with patch('engine.jobs.extract_frames', side_effect=RuntimeError('Fallo temporal de prueba')):
                state = advance_job(store, job_id)
            self.assertEqual(state['status'], 'failed')
            self.assertEqual(state['next_segment'], 1)
            state = advance_job(store, job_id, retry=True)
            while state['status'] not in TERMINAL | {'failed'}:
                state = advance_job(store, job_id)
            self.assertEqual(state['status'], 'extracted_without_ai', state.get('error'))
            self.assertEqual(checkpoint, store.get('analysis', f'{job_id}/segment_000000/result.json'))

    def test_native_editable_annotations_and_complete_map(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            data = json.loads((ROOT / 'example_steps.json').read_text())
            data['steps'] = [dict(data['steps'][0], title=f'Paso {i}') for i in range(12)]
            file = presentation(data, output_dir=directory)
            deck = Presentation(file)
            self.assertEqual(len(deck.slides), 1 + 3 + 12 + 1)
            slide = deck.slides[4]
            xml = slide._element.xml
            self.assertIn('<p:pic>', xml)
            self.assertIn('<p:cxnSp>', xml)
            self.assertIn('type="triangle"', xml)
            self.assertIn('Instrucción completa', slide.notes_slide.notes_text_frame.text)
            # Overlay is a native rectangle and text, not flattened into the picture.
            self.assertIn('prst="rect"', xml)
            self.assertGreater(sum(1 for s in slide.shapes if s.has_text_frame), 4)


if __name__ == '__main__':
    unittest.main()
